from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "data-relationships.md"
OUT_DIR = ROOT / "docs" / "generated"
DOCX_OUT = OUT_DIR / "ERP頁面資料關係說明.docx"
IMG_DIR = OUT_DIR / "images"

FONT_REGULAR = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
FONT_FALLBACK = "/System/Library/Fonts/STHeiti Medium.ttc"
DOC_FONT = "Arial"
CJK_FONT = "Arial Unicode MS"


@dataclass
class MdTable:
    headers: list[str]
    rows: list[list[str]]


def font(size: int) -> ImageFont.FreeTypeFont:
    path = FONT_REGULAR if Path(FONT_REGULAR).exists() else FONT_FALLBACK
    return ImageFont.truetype(path, size)


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, width: int) -> list[str]:
    text = text.replace(" / ", "/").replace("、", "、")
    lines: list[str] = []
    for part in text.split("\n"):
        current = ""
        for ch in part:
            trial = current + ch
            if draw.textbbox((0, 0), trial, font=fnt)[2] <= width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = ch
        lines.append(current)
    return lines


def box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    body: list[str],
    fill: str,
    outline: str = "#8A8F98",
) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=2)
    title_font = font(25)
    body_font = font(19)
    draw.text((x1 + 18, y1 + 14), title, font=title_font, fill="#111111")
    y = y1 + 52
    for line in body:
        for wrapped in wrap(draw, line, body_font, x2 - x1 - 36):
            draw.text((x1 + 18, y), wrapped, font=body_font, fill="#333333")
            y += 27


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], label: str = "") -> None:
    sx, sy = start
    ex, ey = end
    draw.line((sx, sy, ex, ey), fill="#555555", width=4)
    dx = ex - sx
    dy = ey - sy
    if abs(dx) >= abs(dy):
        head = [(ex, ey), (ex - 14 if dx > 0 else ex + 14, ey - 9), (ex - 14 if dx > 0 else ex + 14, ey + 9)]
    else:
        head = [(ex, ey), (ex - 9, ey - 14 if dy > 0 else ey + 14), (ex + 9, ey - 14 if dy > 0 else ey + 14)]
    draw.polygon(head, fill="#555555")
    if label:
        lbl_font = font(17)
        mx, my = (sx + ex) // 2, (sy + ey) // 2
        tw = draw.textbbox((0, 0), label, font=lbl_font)[2]
        draw.rectangle((mx - tw // 2 - 8, my - 18, mx + tw // 2 + 8, my + 10), fill="#FFFFFF")
        draw.text((mx - tw // 2, my - 15), label, font=lbl_font, fill="#333333")


def draw_module_overview(path: Path) -> None:
    img = Image.new("RGB", (1800, 1000), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    box(draw, (70, 110, 400, 310), "主檔資料", ["進銷項目 Item", "人員管理 Employee", "活動管理 Event"], "#F4F8FF")
    box(draw, (500, 80, 850, 330), "營收登記", ["Sale 營收主檔", "SaleLine 銷售明細", "付款方式 / 狀態 / 日期"], "#FFF7E8")
    box(draw, (980, 80, 1320, 330), "營收輸出", ["數據總覽", "月報表", "年報表"], "#F0FFF4")
    box(draw, (500, 430, 850, 650), "庫存紀錄", ["InventoryMovement", "銷售扣庫存", "進貨 / 調整 / 耗用"], "#F7F2FF")
    box(draw, (980, 430, 1320, 650), "費用與人員", ["固定支出", "零用支出", "排班紀錄"], "#F9FAFB")
    box(draw, (1420, 260, 1710, 520), "權限設定", ["User", "Role", "pagePermissions"], "#F8F8F8")
    arrow(draw, (400, 210), (500, 210), "選品項/人員")
    arrow(draw, (850, 200), (980, 200), "Sale/SaleLine")
    arrow(draw, (675, 330), (675, 430), "需庫存")
    arrow(draw, (850, 540), (980, 540), "成本/支出")
    arrow(draw, (1320, 300), (1420, 390), "頁面/API")
    img.save(path)


def draw_revenue_core(path: Path) -> None:
    img = Image.new("RGB", (1800, 930), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    box(draw, (70, 110, 400, 310), "進銷項目", ["Item.id", "Item.type / cost", "requiresInventory"], "#F4F8FF")
    box(draw, (70, 390, 400, 590), "人員管理", ["Employee.id", "銷售人員", "服務人員"], "#F4F8FF")
    box(draw, (520, 120, 850, 580), "營收登記", ["Sale", "id / soldAt / status", "paymentMethod / total", "", "SaleLine", "saleId / itemId", "quantity / lineTotal"], "#FFF7E8")
    box(draw, (980, 90, 1320, 270), "數據總覽", ["Sale.total", "SaleLine.lineTotal", "今日 / 本月 / 品項"], "#F0FFF4")
    box(draw, (980, 330, 1320, 520), "月報表", ["營收 / 毛利", "付款方式", "人員給付"], "#F0FFF4")
    box(draw, (980, 590, 1320, 760), "年報表", ["年度營收", "年度成本", "年度利潤"], "#F0FFF4")
    box(draw, (1420, 330, 1720, 560), "庫存紀錄", ["InventoryMovement", "itemId", "type=SALE", "quantity<0"], "#F7F2FF")
    arrow(draw, (400, 210), (520, 250), "itemId")
    arrow(draw, (400, 490), (520, 440), "employeeId")
    arrow(draw, (850, 220), (980, 180), "營收彙總")
    arrow(draw, (850, 350), (980, 425), "月彙總")
    arrow(draw, (850, 490), (980, 675), "年彙總")
    arrow(draw, (850, 350), (1420, 440), "銷售扣庫存")
    img.save(path)


def draw_field_links(path: Path) -> None:
    img = Image.new("RGB", (1800, 1050), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    nodes = {
        "Item": ((80, 90, 400, 290), ["id", "type", "cost", "eventId", "servicePersonId"]),
        "Sale": ((620, 90, 940, 330), ["id", "soldAt", "status", "salesPersonId", "servicePersonId", "total"]),
        "SaleLine": ((620, 430, 940, 650), ["saleId", "itemId", "quantity", "lineTotal"]),
        "Employee": ((80, 410, 400, 610), ["id", "name", "tags"]),
        "Inventory": ((1120, 430, 1480, 650), ["itemId", "type", "quantity", "reason"]),
        "Reports": ((1120, 90, 1480, 330), ["Sale.soldAt", "Sale.status", "SaleLine.lineTotal", "Item.cost"]),
        "Event": ((80, 730, 400, 930), ["id", "title", "startsAt", "endsAt"]),
        "Registration": ((620, 730, 940, 970), ["eventId", "itemId", "saleId", "paidAmount"]),
    }
    fills = {"Sale": "#FFF7E8", "SaleLine": "#FFF7E8", "Reports": "#F0FFF4", "Inventory": "#F7F2FF"}
    for name, (xy, body) in nodes.items():
        box(draw, xy, name, body, fills.get(name, "#F4F8FF"))
    arrow(draw, (400, 190), (620, 540), "Item.id = SaleLine.itemId")
    arrow(draw, (780, 330), (780, 430), "Sale.id = SaleLine.saleId")
    arrow(draw, (400, 510), (620, 210), "Employee.id = Sale.*PersonId")
    arrow(draw, (940, 540), (1120, 540), "SaleLine -> 扣庫存")
    arrow(draw, (940, 210), (1120, 210), "Sale/SaleLine -> 報表")
    arrow(draw, (400, 830), (620, 850), "Event.id")
    arrow(draw, (400, 190), (620, 900), "Item.id")
    arrow(draw, (780, 730), (780, 650), "付款後 Sale.id")
    img.save(path)


def clean_inline(text: str) -> str:
    return text.replace("`", "").replace("->", "→")


def set_run_font(run, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = DOC_FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_runs(paragraph, size: int | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size)


def parse_table(lines: list[str], start: int) -> tuple[MdTable, int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].startswith("|"):
        cells = [clean_inline(c.strip()) for c in lines[i].strip().strip("|").split("|")]
        rows.append(cells)
        i += 1
    headers = rows[0]
    body = rows[2:] if len(rows) > 1 and set(rows[1][0].replace("-", "").replace(" ", "")) == set() else rows[1:]
    return MdTable(headers, body), i


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "請在 Word 或 Google Docs 更新目錄"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, placeholder, fld_end])


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_doc_table(doc: Document, md_table: MdTable, small: bool = False) -> None:
    table = doc.add_table(rows=1, cols=len(md_table.headers))
    table.style = "Table Grid"
    table.autofit = True
    header = table.rows[0].cells
    for idx, value in enumerate(md_table.headers):
        header[idx].text = value
        set_cell_shading(header[idx], "F1F3F4")
        for p in header[idx].paragraphs:
            for r in p.runs:
                set_run_font(r, size=8 if small else 9, bold=True)
    for row in md_table.rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row[: len(md_table.headers)]):
            cells[idx].text = value
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_run_font(r, size=7 if small else 8)
    doc.add_paragraph()


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.orientation = WD_ORIENT.PORTRAIT

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = DOC_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after, color in [
        ("Heading 1", 20, 20, 6, "000000"),
        ("Heading 2", 16, 18, 6, "000000"),
        ("Heading 3", 14, 16, 4, "434343"),
    ]:
        st = styles[name]
        st.font.name = DOC_FONT
        st._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        st.font.size = Pt(size)
        st.font.bold = False
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def add_manual_toc(doc: Document) -> None:
    p = doc.add_heading("", level=1)
    set_run_font(p.add_run("目錄"), size=20)
    entries = [
        "一、讀圖方式",
        "二、頁面模組總覽",
        "三、營收核心關聯",
        "四、模組、頁面與資料庫關聯",
        "五、頁面欄位層級關聯",
        "六、核心資料流",
        "七、報表計算資料來源",
    ]
    for entry in entries:
        p = doc.add_paragraph(style="List Number")
        set_run_font(p.add_run(entry))
    p = doc.add_paragraph()
    add_toc_field(p)
    doc.add_page_break()


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    module_img = IMG_DIR / "module-overview.png"
    revenue_img = IMG_DIR / "revenue-core.png"
    field_img = IMG_DIR / "field-links.png"
    draw_module_overview(module_img)
    draw_revenue_core(revenue_img)
    draw_field_links(field_img)

    doc = Document()
    configure_styles(doc)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("ERP 頁面資料關係說明")
    set_run_font(title_run, size=26, bold=False)
    subtitle = doc.add_paragraph()
    set_run_font(subtitle.add_run("UNV.PLAN ERP｜資料庫、頁面與欄位層級關聯整理"))
    add_manual_toc(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    skip_mermaid = False
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("```mermaid"):
            skip_mermaid = True
            i += 1
            continue
        if skip_mermaid:
            if line.startswith("```"):
                skip_mermaid = False
            i += 1
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            p = doc.add_heading("", level=1)
            set_run_font(p.add_run(heading), size=20)
            if heading.startswith("二、"):
                doc.add_picture(str(module_img), width=Inches(6.5))
            if heading.startswith("三、"):
                doc.add_picture(str(revenue_img), width=Inches(6.5))
            if heading.startswith("五、"):
                doc.add_picture(str(field_img), width=Inches(6.5))
            i += 1
            continue
        if line.startswith("### "):
            heading = line[4:].strip()
            if heading in {"精準欄位連線圖", "資料表關聯圖"}:
                i += 1
                continue
            p = doc.add_heading("", level=2)
            set_run_font(p.add_run(heading), size=16)
            i += 1
            continue
        if line.startswith("|"):
            table, i = parse_table(lines, i)
            add_doc_table(doc, table, small=len(table.headers) >= 6)
            continue
        if not line.strip():
            i += 1
            continue
        paragraph = doc.add_paragraph()
        set_run_font(paragraph.add_run(clean_inline(line.strip())))
        if re.match(r"^[0-9]+\\.", line.strip()):
            paragraph.style = "List Number"
        i += 1

    doc.save(DOCX_OUT)


if __name__ == "__main__":
    build()
