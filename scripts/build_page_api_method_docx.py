from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path("docs/page-api-method-list.docx")


ROWS = [
    ("/", "首頁轉址", "-", "-", "直接轉到 /dashboard；沒有頁面 API 呼叫。"),
    ("/login", "登入 ERP", "POST", "/api/auth/login", "提交登入表單。"),
    ("/register", "建立帳號", "POST", "/api/auth/register", "提交註冊表單。"),
    ("/dashboard", "營運總覽 / 數據總覽", "SERVER", "Prisma direct query", "初始今日/月營收與品項占比在 server component 直接查 Prisma。"),
    ("/dashboard", "營運總覽 / 數據總覽", "GET", "/api/reports/daily?date=YYYY-MM-DD", "切換日期後讀取日營收。"),
    ("/sales", "營收登記", "GET", "/api/items", "讀取可銷售品項。"),
    ("/sales", "營收登記", "GET", "/api/employees", "讀取銷售/服務人員。"),
    ("/sales", "營收登記", "GET", "/api/sales", "讀取營收紀錄。"),
    ("/sales", "營收登記", "POST", "/api/sales", "新增營收。"),
    ("/sales", "營收登記", "PATCH", "/api/sales/:id", "編輯營收。"),
    ("/sales", "營收登記", "DELETE", "/api/sales/:id", "刪除營收。"),
    ("/permissions", "權限設定", "GET", "/api/permissions", "讀取權限、角色與使用者。"),
    ("/permissions", "權限設定", "POST", "/api/permissions", "新增角色或帳號。"),
    ("/permissions", "權限設定", "PATCH", "/api/permissions/users/:id", "更新帳號資料或啟用狀態。"),
    ("/permissions", "權限設定", "PATCH", "/api/permissions/roles/:id", "更新角色資料或啟用狀態。"),
    ("/reports/daily", "日報表", "GET", "/api/reports/daily?date=YYYY-MM-DD", "讀取指定日期報表。"),
    ("/reports/monthly", "月報表", "GET", "/api/reports/monthly?date=YYYY-MM-DD", "讀取指定月份報表。"),
    ("/reports/yearly", "年報表 / 股東利潤", "GET", "/api/reports/yearly?date=YYYY-MM-DD", "讀取指定年度報表。"),
    ("/items", "進銷項目", "GET", "/api/items", "讀取進銷項目。"),
    ("/items", "進銷項目", "POST", "/api/items", "新增進銷項目。"),
    ("/items", "進銷項目", "PATCH", "/api/items/:id", "編輯進銷項目。"),
    ("/items", "進銷項目", "DELETE", "/api/items/:id", "刪除進銷項目。"),
    ("/items", "進銷項目", "POST", "/api/inventory", "批次新增需要庫存的品項時建立開帳庫存。"),
    ("/inventory", "庫存紀錄", "GET", "/api/inventory", "讀取庫存紀錄。"),
    ("/inventory", "庫存紀錄", "POST", "/api/inventory", "新增庫存異動。"),
    ("/inventory", "庫存紀錄", "PATCH", "/api/inventory/:id", "編輯庫存異動。"),
    ("/inventory", "庫存紀錄", "DELETE", "/api/inventory/:id", "刪除庫存異動。"),
    ("/inventory", "庫存紀錄", "GET", "/api/items", "讀取品項下拉與目前庫存。"),
    ("/assets", "固定/變動資產", "GET", "/api/assets", "讀取資產。"),
    ("/assets", "固定/變動資產", "POST", "/api/assets", "新增資產。"),
    ("/assets", "固定/變動資產", "PATCH", "/api/assets/:id", "編輯資產。"),
    ("/assets", "固定/變動資產", "DELETE", "/api/assets/:id", "刪除或停用資產。"),
    ("/assets", "固定/變動資產", "POST", "/api/assets/copy-fixed-month", "複製固定資產到指定月份。"),
    ("/petty-cash", "零用金 / 零用支出", "GET", "/api/petty-cash", "讀取零用金紀錄。"),
    ("/petty-cash", "零用金 / 零用支出", "POST", "/api/petty-cash", "新增零用金紀錄。"),
    ("/petty-cash", "零用金 / 零用支出", "PATCH", "/api/petty-cash/:id", "編輯紀錄或更新出帳狀態。"),
    ("/petty-cash", "零用金 / 零用支出", "DELETE", "/api/petty-cash/:id", "刪除零用金紀錄。"),
    ("/petty-cash", "零用金 / 零用支出", "GET", "/api/employees", "讀取代墊者下拉。"),
    ("/fixed-expenses", "固定支出", "GET", "/api/fixed-expenses", "讀取固定支出。"),
    ("/fixed-expenses", "固定支出", "POST", "/api/fixed-expenses", "新增固定支出。"),
    ("/fixed-expenses", "固定支出", "PATCH", "/api/fixed-expenses/:id", "編輯固定支出。"),
    ("/fixed-expenses", "固定支出", "DELETE", "/api/fixed-expenses/:id", "刪除固定支出。"),
    ("/fixed-expenses", "固定支出", "GET", "/api/employees", "讀取支出對象下拉。"),
    ("/fixed-expenses", "固定支出", "GET", "/api/shifts", "讀取排班，用於固定支出相關統計。"),
    ("/employees", "人員管理", "GET", "/api/employees", "讀取人員。"),
    ("/employees", "人員管理", "POST", "/api/employees", "新增人員。"),
    ("/employees", "人員管理", "PATCH", "/api/employees/:id", "編輯人員。"),
    ("/employees", "人員管理", "DELETE", "/api/employees/:id", "刪除或停用人員。"),
    ("/employees", "人員管理", "GET", "/api/employee-tags", "讀取人員 TAG。"),
    ("/employees", "人員管理", "POST", "/api/employee-tags", "新增人員 TAG。"),
    ("/employees", "人員管理", "DELETE", "/api/employee-tags/:id", "刪除人員 TAG。"),
    ("/shifts", "排班紀錄", "GET", "/api/shifts", "讀取排班與備註。"),
    ("/shifts", "排班紀錄", "GET", "/api/employees", "讀取可排班人員。"),
    ("/shifts", "排班紀錄", "POST", "/api/shifts", "新增單日排班或區間備註。"),
    ("/shifts", "排班紀錄", "DELETE", "/api/shifts/:id", "刪除排班或備註。"),
    ("/sops", "工作流程 SOP / SOP", "GET", "/api/sops", "讀取 SOP。"),
    ("/sops", "工作流程 SOP / SOP", "POST", "/api/sops", "新增 SOP。"),
    ("/sops", "工作流程 SOP / SOP", "PATCH", "/api/sops/:id", "編輯 SOP。"),
    ("/sops", "工作流程 SOP / SOP", "DELETE", "/api/sops/:id", "刪除 SOP。"),
    ("/sop-steps", "SOP 步驟", "GET", "/api/sop-steps", "讀取 SOP 步驟。"),
    ("/sop-steps", "SOP 步驟", "POST", "/api/sop-steps", "新增 SOP 步驟。"),
    ("/sop-steps", "SOP 步驟", "PATCH", "/api/sop-steps/:id", "編輯 SOP 步驟。"),
    ("/sop-steps", "SOP 步驟", "DELETE", "/api/sop-steps/:id", "刪除 SOP 步驟。"),
    ("/sop-steps", "SOP 步驟", "GET", "/api/sops", "讀取 SOP 下拉。"),
    ("/events", "活動管理", "GET", "/api/events", "活動分頁讀取活動。"),
    ("/events", "活動管理", "POST", "/api/events", "新增活動，含活動品項批次資料。"),
    ("/events", "活動管理", "PATCH", "/api/events/:id", "編輯活動。"),
    ("/events", "活動管理", "DELETE", "/api/events/:id", "刪除活動。"),
    ("/events", "活動管理", "GET", "/api/employees", "活動品項服務人員下拉。"),
    ("/events", "活動管理", "GET", "/api/items", "活動品項同步後重新讀取品項。"),
    ("/events", "活動報名分頁", "GET", "/api/event-registrations", "報名分頁讀取報名。"),
    ("/events", "活動報名分頁", "POST", "/api/event-registrations", "新增活動報名。"),
    ("/events", "活動報名分頁", "PATCH", "/api/event-registrations/:id", "編輯報名或付款狀態。"),
    ("/events", "活動報名分頁", "DELETE", "/api/event-registrations/:id", "刪除報名。"),
    ("/events", "活動報名分頁", "GET", "/api/events", "報名表單活動下拉。"),
    ("/events", "活動報名分頁", "GET", "/api/items", "報名表單品項下拉。"),
    ("/event-registrations", "活動報名登記", "GET", "/api/event-registrations", "讀取活動報名。"),
    ("/event-registrations", "活動報名登記", "POST", "/api/event-registrations", "新增活動報名。"),
    ("/event-registrations", "活動報名登記", "PATCH", "/api/event-registrations/:id", "編輯報名或付款狀態。"),
    ("/event-registrations", "活動報名登記", "DELETE", "/api/event-registrations/:id", "刪除活動報名。"),
    ("/event-registrations", "活動報名登記", "GET", "/api/events", "活動下拉。"),
    ("/event-registrations", "活動報名登記", "GET", "/api/items", "品項下拉。"),
    ("/customers", "客服資訊", "GET", "/api/customers", "讀取客戶資料。"),
    ("/customers", "客服資訊", "POST", "/api/customers", "新增客戶資料。"),
    ("/customers", "客服資訊", "PATCH", "/api/customers/:id", "編輯客戶資料。"),
    ("/customers", "客服資訊", "DELETE", "/api/customers/:id", "刪除客戶資料。"),
    ("/customer-service-records", "客服使用紀錄 / 客服紀錄", "GET", "/api/customer-service-records", "讀取客服使用紀錄。"),
    ("/customer-service-records", "客服使用紀錄 / 客服紀錄", "POST", "/api/customer-service-records", "新增客服使用紀錄。"),
    ("/customer-service-records", "客服使用紀錄 / 客服紀錄", "PATCH", "/api/customer-service-records/:id", "編輯客服使用紀錄。"),
    ("/customer-service-records", "客服使用紀錄 / 客服紀錄", "DELETE", "/api/customer-service-records/:id", "刪除客服使用紀錄。"),
    ("/customer-service-records", "客服使用紀錄 / 客服紀錄", "GET", "/api/customers", "客戶下拉。"),
    ("/course-payments", "線上課程收費", "GET", "/api/course-payments", "讀取課程收費紀錄。"),
    ("/course-payments", "線上課程收費", "POST", "/api/course-payments", "新增課程收費紀錄。"),
    ("/course-payments", "線上課程收費", "PATCH", "/api/course-payments/:id", "編輯課程收費紀錄。"),
    ("/course-payments", "線上課程收費", "DELETE", "/api/course-payments/:id", "刪除課程收費紀錄。"),
    ("/course-payments", "線上課程收費", "GET", "/api/customers", "客戶下拉。"),
]


VISIBLE_ROWS = [
    ("/login", "登入 ERP", "POST", "/api/auth/login", "提交登入表單。"),
    ("/register", "建立帳號", "POST", "/api/auth/register", "提交註冊表單。"),
    ("/dashboard", "營運總覽 / 數據總覽", "GET", "/api/reports/daily?date=YYYY-MM-DD", "切換日期後讀取日營收。初始數據總覽為 server component 直接查 Prisma，未列為 API。"),
    ("/sales", "營收登記", "GET", "/api/items", "讀取可銷售品項。"),
    ("/sales", "營收登記", "GET", "/api/employees", "讀取銷售/服務人員。"),
    ("/sales", "營收登記", "GET", "/api/sales", "讀取營收紀錄。"),
    ("/sales", "營收登記", "POST", "/api/sales", "新增營收。"),
    ("/sales", "營收登記", "PATCH", "/api/sales/:id", "編輯營收。"),
    ("/sales", "營收登記", "DELETE", "/api/sales/:id", "刪除營收。"),
    ("/items", "進銷項目", "GET", "/api/items", "讀取進銷項目。"),
    ("/items", "進銷項目", "POST", "/api/items", "新增進銷項目。"),
    ("/items", "進銷項目", "PATCH", "/api/items/:id", "編輯進銷項目。"),
    ("/items", "進銷項目", "DELETE", "/api/items/:id", "刪除進銷項目。"),
    ("/items", "進銷項目", "POST", "/api/inventory", "批次新增需要庫存的品項時建立開帳庫存。"),
    ("/employees", "人員管理", "GET", "/api/employees", "讀取人員。"),
    ("/employees", "人員管理", "POST", "/api/employees", "新增人員。"),
    ("/employees", "人員管理", "PATCH", "/api/employees/:id", "編輯人員。"),
    ("/employees", "人員管理", "DELETE", "/api/employees/:id", "刪除或停用人員。"),
    ("/employees", "人員管理", "GET", "/api/employee-tags", "讀取人員 TAG。"),
    ("/employees", "人員管理", "POST", "/api/employee-tags", "新增人員 TAG。"),
    ("/employees", "人員管理", "DELETE", "/api/employee-tags/:id", "刪除人員 TAG。"),
    ("/shifts", "排班紀錄", "GET", "/api/shifts", "讀取排班與備註。"),
    ("/shifts", "排班紀錄", "GET", "/api/employees", "讀取可排班人員。"),
    ("/shifts", "排班紀錄", "POST", "/api/shifts", "新增單日排班或區間備註。"),
    ("/shifts", "排班紀錄", "DELETE", "/api/shifts/:id", "刪除排班或備註。"),
    ("/reports/monthly", "月報表", "GET", "/api/reports/monthly?date=YYYY-MM-DD", "讀取指定月份報表。"),
    ("/reports/yearly", "年報表 / 股東利潤", "GET", "/api/reports/yearly?date=YYYY-MM-DD", "讀取指定年度報表。"),
    ("/inventory", "庫存紀錄", "GET", "/api/inventory", "讀取庫存紀錄。"),
    ("/inventory", "庫存紀錄", "POST", "/api/inventory", "新增庫存異動。"),
    ("/inventory", "庫存紀錄", "PATCH", "/api/inventory/:id", "編輯庫存異動。"),
    ("/inventory", "庫存紀錄", "DELETE", "/api/inventory/:id", "刪除庫存異動。"),
    ("/inventory", "庫存紀錄", "GET", "/api/items", "讀取品項下拉與目前庫存。"),
    ("/petty-cash", "零用金 / 零用支出", "GET", "/api/petty-cash", "讀取零用金紀錄。"),
    ("/petty-cash", "零用金 / 零用支出", "POST", "/api/petty-cash", "新增零用金紀錄。"),
    ("/petty-cash", "零用金 / 零用支出", "PATCH", "/api/petty-cash/:id", "編輯紀錄或更新出帳狀態。"),
    ("/petty-cash", "零用金 / 零用支出", "DELETE", "/api/petty-cash/:id", "刪除零用金紀錄。"),
    ("/petty-cash", "零用金 / 零用支出", "GET", "/api/employees", "讀取代墊者下拉。"),
    ("/fixed-expenses", "固定支出", "GET", "/api/fixed-expenses", "讀取固定支出。"),
    ("/fixed-expenses", "固定支出", "POST", "/api/fixed-expenses", "新增固定支出。"),
    ("/fixed-expenses", "固定支出", "PATCH", "/api/fixed-expenses/:id", "編輯固定支出。"),
    ("/fixed-expenses", "固定支出", "DELETE", "/api/fixed-expenses/:id", "刪除固定支出。"),
    ("/fixed-expenses", "固定支出", "GET", "/api/employees", "讀取支出對象下拉。"),
    ("/fixed-expenses", "固定支出", "GET", "/api/shifts", "讀取排班，用於固定支出相關統計。"),
    ("/events", "活動管理", "GET", "/api/events", "活動分頁讀取活動。"),
    ("/events", "活動管理", "POST", "/api/events", "新增活動，含活動品項批次資料。"),
    ("/events", "活動管理", "PATCH", "/api/events/:id", "編輯活動。"),
    ("/events", "活動管理", "DELETE", "/api/events/:id", "刪除活動。"),
    ("/events", "活動管理", "GET", "/api/employees", "活動品項服務人員下拉。"),
    ("/events", "活動管理", "GET", "/api/items", "活動品項同步後重新讀取品項。"),
    ("/events", "活動報名分頁", "GET", "/api/event-registrations", "報名分頁讀取報名。"),
    ("/events", "活動報名分頁", "POST", "/api/event-registrations", "新增活動報名。"),
    ("/events", "活動報名分頁", "PATCH", "/api/event-registrations/:id", "編輯報名或付款狀態。"),
    ("/events", "活動報名分頁", "DELETE", "/api/event-registrations/:id", "刪除報名。"),
    ("/events", "活動報名分頁", "GET", "/api/events", "報名表單活動下拉。"),
    ("/events", "活動報名分頁", "GET", "/api/items", "報名表單品項下拉。"),
    ("/sops", "工作流程 SOP / SOP", "GET", "/api/sops", "讀取 SOP。"),
    ("/sops", "工作流程 SOP / SOP", "POST", "/api/sops", "新增 SOP。"),
    ("/sops", "工作流程 SOP / SOP", "PATCH", "/api/sops/:id", "編輯 SOP。"),
    ("/sops", "工作流程 SOP / SOP", "DELETE", "/api/sops/:id", "刪除 SOP。"),
    ("/permissions", "權限設定", "GET", "/api/permissions", "讀取權限、角色與使用者。"),
    ("/permissions", "權限設定", "POST", "/api/permissions", "新增角色或帳號。"),
    ("/permissions", "權限設定", "PATCH", "/api/permissions/users/:id", "更新帳號資料或啟用狀態。"),
    ("/permissions", "權限設定", "PATCH", "/api/permissions/roles/:id", "更新角色資料或啟用狀態。"),
]


METHOD_FILL = {
    "GET": "EAF4F7",
    "POST": "EAF5EA",
    "PATCH": "FFF3D6",
    "DELETE": "FCE8E6",
    "SERVER": "EFEAF7",
    "-": "F2F4F7",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width(cell, width)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(table, top=80, bottom=80, start=120, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for margin_name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tbl_cell_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_paragraph(paragraph, font_size=9.0, bold=False, color="000000") -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        run.font.name = "PingFang TC"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def add_cell_text(cell, text: str, *, bold=False, font_size=9.0, color="000000", align=None) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    run = paragraph.add_run(text)
    if align is not None:
        paragraph.alignment = align
    style_paragraph(paragraph, font_size=font_size, bold=bold, color=color)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "PingFang TC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = "PingFang TC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_footer(doc: Document) -> None:
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("UNV.PLAN ERP API Reference")
    run.font.name = "PingFang TC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("666666")


def build_docx() -> None:
    doc = Document()
    configure_styles(doc)
    add_footer(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title_run = title.add_run("ERP 頁面 API 對照表")
    title_run.font.name = "PingFang TC"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle_run = subtitle.add_run("依頁面中文名稱與 HTTP method 分列列出。來源：目前 Next.js 頁面與 React 元件 fetch 呼叫。")
    subtitle_run.font.name = "PingFang TC"
    subtitle_run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")
    subtitle_run.font.size = Pt(9.5)
    subtitle_run.font.color.rgb = RGBColor.from_string("555555")

    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(8)
    note_run = note.add_run("備註：SERVER 代表該頁面在 server component 直接查 Prisma，不是瀏覽器端 fetch API。")
    note_run.font.name = "PingFang TC"
    note_run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = RGBColor.from_string("555555")

    headers = ["頁面", "中文名稱", "Method", "API", "用途 / 備註"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"

    widths = [1400, 1700, 1080, 2760, 2420]
    set_table_geometry(table, widths)
    set_cell_margins(table)

    header_row = table.rows[0]
    header_row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
    set_repeat_table_header(header_row)
    for cell, header in zip(header_row.cells, headers):
        set_cell_shading(cell, "E8EEF5")
        add_cell_text(cell, header, bold=True, font_size=9.5, color="0B2545", align=WD_ALIGN_PARAGRAPH.CENTER)

    for page, chinese_name, method, api, note_text in VISIBLE_ROWS:
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        values = [page, chinese_name, method, api, note_text]
        for idx, (cell, value) in enumerate(zip(row.cells, values)):
            if idx == 2:
                set_cell_shading(cell, METHOD_FILL.get(method, "FFFFFF"))
                add_cell_text(cell, value, bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            elif idx == 0:
                add_cell_text(cell, value, font_size=8.5)
            elif idx == 3:
                add_cell_text(cell, value, font_size=8.5)
            else:
                add_cell_text(cell, value, font_size=8.5)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_docx()
