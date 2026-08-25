from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path("docs/erp-page-api-reference.docx")
FONT = "PingFang TC"


SECTIONS = [
    {
        "module": "認證",
        "pages": [
            {
                "page": "/login",
                "name": "登入 ERP",
                "apis": [
                    {
                        "method": "POST",
                        "api": "/api/auth/login",
                        "content": "提交登入表單，驗證帳號密碼後設定 session cookie。",
                        "fields": "email：Email；password：密碼。",
                        "rules": "帳號必須 active=true；密碼需正確；非 Admin 帳號的角色也必須 active=true；失敗回 401。"
                    }
                ]
            },
            {
                "page": "/register",
                "name": "建立帳號",
                "apis": [
                    {
                        "method": "POST",
                        "api": "/api/auth/register",
                        "content": "建立第一個本機 Admin 帳號並登入。",
                        "fields": "name：姓名；email：Email；password：密碼。",
                        "rules": "姓名與 Email 必填；密碼至少 6 碼；系統已有帳號後不可從註冊頁新增，需改由權限設定建立。"
                    }
                ]
            }
        ]
    },
    {
        "module": "營收",
        "pages": [
            {
                "page": "/sales",
                "name": "營收登記",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/items",
                        "content": "讀取可銷售品項，下拉使用。",
                        "fields": "Item：id、type、name、price、cost、requiresInventory、active、notes、eventId、servicePersonId、stockQuantity。",
                        "rules": "需 items 權限；前端會過濾 active=false；stockQuantity 由庫存異動加總。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/employees",
                        "content": "讀取銷售人員與服務人員下拉。",
                        "fields": "Employee：id、name、role、active、tags。",
                        "rules": "需 employees 權限；前端會過濾 active=false；服務人員依標籤內部員工/外部員工篩選。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/sales",
                        "content": "讀取營收紀錄，含銷售人員、服務人員與明細品項。",
                        "fields": "Sale：id、soldAt、customerName、paymentMethod、notes、subtotal、discount、total、status、salesPerson、servicePerson、lines。",
                        "rules": "需 sales 權限；依 soldAt desc 排序。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/sales",
                        "content": "新增營收；目前每筆銷售明細會各自建立一張 Sale。",
                        "fields": "soldAt、customerName、salesPersonId、paymentMethod、status、notes、lines[{itemId, servicePersonId, quantity, unitPrice}]。",
                        "rules": "salesPersonId 必填；至少一筆有效 lines；品項必須 active=true；非香油捐贈品項必填 servicePersonId；requiresInventory 品項會新增 SALE 庫存異動且 quantity 為負數。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/sales/:id",
                        "content": "更新營收與明細。",
                        "fields": "soldAt、customerName、salesPersonId、paymentMethod、status、notes、discount、lines[{itemId, servicePersonId, quantity, unitPrice}]。",
                        "rules": "會先刪除原 SaleLine 與 reason=銷售單 {id} 的庫存異動，再重建明細與庫存扣減；discount 從 subtotal 扣除，total 不小於 0。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/sales/:id",
                        "content": "刪除營收單。",
                        "fields": "路徑 id。",
                        "rules": "會刪除 reason=銷售單 {id} 的庫存異動，再刪 Sale；SaleLine 由 cascade 刪除。"
                    }
                ]
            },
            {
                "page": "/dashboard",
                "name": "數據總覽",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/reports/daily?date=YYYY-MM-DD",
                        "content": "切換日營收日期時讀取日報表摘要。",
                        "fields": "query.date：基準日期；回傳 summary.revenue 等報表欄位。",
                        "rules": "初始數據總覽不是 API，為 server component 直接查 Prisma；此處只列切換日期時的 API。"
                    }
                ]
            }
        ]
    },
    {
        "module": "管理",
        "pages": [
            {
                "page": "/items",
                "name": "進銷項目",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/items",
                        "content": "讀取所有進銷項目並附加目前庫存數。",
                        "fields": "id、type、name、price、cost、requiresInventory、active、notes、eventId、servicePersonId、createdAt、updatedAt、stockQuantity。",
                        "rules": "需 items 權限；依 createdAt desc 排序；stockQuantity 為 InventoryMovement.quantity 加總。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/items",
                        "content": "新增進銷項目。",
                        "fields": "type、name、price、cost、requiresInventory、notes、active。",
                        "rules": "price/cost 會轉數字且不可小於 0；預設 active=true、cost=0、price=0、requiresInventory=false；type=OIL_DONATION 時強制 requiresInventory=false。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/items/:id",
                        "content": "更新進銷項目。",
                        "fields": "type、name、price、cost、requiresInventory、notes、active；路徑 id 指定品項。",
                        "rules": "套用同 POST 正規化規則。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/items/:id",
                        "content": "刪除進銷項目。",
                        "fields": "路徑 id。",
                        "rules": "若已有 SaleLine 會回 409 不可刪；可刪時會先刪該品項 InventoryMovement，再刪 Item。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/inventory",
                        "content": "批次新增需要庫存的品項時建立開帳庫存。",
                        "fields": "itemId、type=PURCHASE、quantity、reason。",
                        "rules": "由進銷項目頁的批次新增流程觸發；quantity 必須大於 0。"
                    }
                ]
            },
            {
                "page": "/employees",
                "name": "人員管理",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/employees",
                        "content": "讀取人員與 TAG。",
                        "fields": "id、name、role、active、tags、createdAt、updatedAt。",
                        "rules": "需 employees 權限；soft delete 模組，刪除會改 active=false。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/employees",
                        "content": "新增人員。",
                        "fields": "name、role、tagIds。",
                        "rules": "預設 active=true；tagIds 會轉為 tags connect。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/employees/:id",
                        "content": "更新人員。",
                        "fields": "name、role、tagIds。",
                        "rules": "tagIds 會覆蓋原 tags set。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/employees/:id",
                        "content": "停用人員。",
                        "fields": "路徑 id。",
                        "rules": "soft delete：更新 active=false，不實體刪除。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/employee-tags",
                        "content": "讀取人員 TAG。",
                        "fields": "id、name。",
                        "rules": "需 employees 權限；依 name asc 排序。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/employee-tags",
                        "content": "新增或取得既有 TAG。",
                        "fields": "name。",
                        "rules": "name 必填；以 name upsert，重複名稱不新增第二筆。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/employee-tags/:id",
                        "content": "刪除 TAG。",
                        "fields": "路徑 id。",
                        "rules": "會先解除 employees 關聯再刪除；失敗回 409。"
                    }
                ]
            },
            {
                "page": "/shifts",
                "name": "排班紀錄",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/shifts",
                        "content": "讀取排班與日曆備註。",
                        "fields": "id、employeeId、type、startsAt、endsAt、location、notes、employee、eventId、createdAt。",
                        "rules": "需 shifts 權限；排班頁同時用 employeeId=null 的紀錄表示日曆備註。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/employees",
                        "content": "讀取可排班員工。",
                        "fields": "Employee 含 tags。",
                        "rules": "前端篩選 active=true 且 TAG 為內部員工/外部員工。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/shifts",
                        "content": "新增單日排班或區間備註。",
                        "fields": "employeeId、type、startsAt、endsAt、location、notes。",
                        "rules": "employeeId 可省略，省略時作為日曆備註；type 為 ONSITE 或 CHARITY。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/shifts/:id",
                        "content": "刪除排班或備註。",
                        "fields": "路徑 id。",
                        "rules": "一般實體刪除；被引用或不存在時可能回 409/錯誤。"
                    }
                ]
            }
        ]
    },
    {
        "module": "報表",
        "pages": [
            {
                "page": "/reports/monthly",
                "name": "月報表",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/reports/monthly?date=YYYY-MM-DD",
                        "content": "讀取指定月份營收、品項、付款方式、活動、支出、人員分潤分析。",
                        "fields": "query.date；回傳 summary、itemAnalysis、itemTypeAnalysis、paymentMethodAnalysis、activityAnalysis、fixedExpenseAnalysis、pettyCashExpenseAnalysis、shiftIncomeAnalysis、serviceCommissionAnalysis、employeePayoutAnalysis。",
                        "rules": "需 reports-monthly 權限；月報另計算上月營收變化率。"
                    }
                ]
            },
            {
                "page": "/reports/yearly",
                "name": "年報表 / 股東利潤",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/reports/yearly?date=YYYY-MM-DD",
                        "content": "讀取指定年度報表資料。",
                        "fields": "query.date；回傳同報表結構。",
                        "rules": "需 reports-yearly 權限；股東分潤比例在前端計算。"
                    }
                ]
            }
        ]
    },
    {
        "module": "資產",
        "pages": [
            {
                "page": "/inventory",
                "name": "庫存紀錄",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/inventory",
                        "content": "讀取庫存異動紀錄並附加該筆異動後庫存。",
                        "fields": "id、itemId、type、quantity、movedAt、reason、item、stockQuantity。",
                        "rules": "需 inventory 權限；stockQuantity 依 itemId、movedAt、createdAt、id 排序累加。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/inventory",
                        "content": "新增庫存異動。",
                        "fields": "itemId、type、quantity、movedAt、reason。",
                        "rules": "type 為 PURCHASE、SALE、ADJUSTMENT、CONSUMPTION；前端對 SALE/CONSUMPTION 會送負數 quantity。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/inventory/:id",
                        "content": "更新庫存異動。",
                        "fields": "itemId、type、quantity、movedAt、reason。",
                        "rules": "movedAt 會轉 Date；quantity 會轉 Number。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/inventory/:id",
                        "content": "刪除庫存異動。",
                        "fields": "路徑 id。",
                        "rules": "一般實體刪除。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/items",
                        "content": "讀取品項下拉與目前庫存。",
                        "fields": "Item 含 stockQuantity。",
                        "rules": "用於選擇品項與顯示目前庫存。"
                    }
                ]
            },
            {
                "page": "/petty-cash",
                "name": "零用支出",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/petty-cash",
                        "content": "讀取零用金紀錄並計算 beforeAmount/afterAmount。",
                        "fields": "id、type、amount、purpose、employeeId、employee、entryDate、checkedOut、beforeAmount、afterAmount。",
                        "rules": "需 petty-cash 權限；餘額依 entryDate、createdAt、id 逐筆累加。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/petty-cash",
                        "content": "新增零用金收入或支出。",
                        "fields": "type、amount、purpose、employeeId、checkedOut、entryDate。",
                        "rules": "type=INCOME 時 employeeId=null 且 checkedOut=false；checkedOut 預設 false。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/petty-cash/:id",
                        "content": "更新零用金紀錄或出帳狀態。",
                        "fields": "type、amount、purpose、employeeId、checkedOut、entryDate；可單獨更新 checkedOut。",
                        "rules": "type=INCOME 時不可設定代墊者與出帳。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/petty-cash/:id",
                        "content": "刪除零用金紀錄。",
                        "fields": "路徑 id。",
                        "rules": "一般實體刪除。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/employees",
                        "content": "讀取代墊者下拉。",
                        "fields": "Employee。",
                        "rules": "前端用於選擇代墊者。"
                    }
                ]
            },
            {
                "page": "/fixed-expenses",
                "name": "固定支出",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/fixed-expenses",
                        "content": "讀取固定支出並附加 cash balance 欄位。",
                        "fields": "id、type、amount、purpose、employeeId、employee、entryDate、beforeAmount、afterAmount。",
                        "rules": "需 fixed-expenses 權限；後端強制 type=EXPENSE。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/fixed-expenses",
                        "content": "新增固定支出。",
                        "fields": "amount、purpose、employeeId、entryDate。",
                        "rules": "purpose 為房租、水電、正職人員薪資、兼職人員薪資、物品修繕；employeeId 必填。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/fixed-expenses/:id",
                        "content": "更新固定支出。",
                        "fields": "amount、purpose、employeeId、entryDate。",
                        "rules": "後端仍強制 type=EXPENSE。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/fixed-expenses/:id",
                        "content": "刪除固定支出。",
                        "fields": "路徑 id。",
                        "rules": "一般實體刪除。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/employees",
                        "content": "讀取支出對象下拉。",
                        "fields": "Employee。",
                        "rules": "用於選擇支出對象。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/shifts",
                        "content": "讀取排班，用於固定支出相關統計。",
                        "fields": "Shift。",
                        "rules": "前端用於固定支出頁統計輔助。"
                    }
                ]
            }
        ]
    },
    {
        "module": "活動",
        "pages": [
            {
                "page": "/events",
                "name": "活動管理 / 活動報名",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/events",
                        "content": "讀取活動與活動同步品項。",
                        "fields": "Event：id、title、startsAt、endsAt、active、items。",
                        "rules": "需 events 權限；活動頁包含活動管理與活動報名兩個分頁。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/events",
                        "content": "新增活動，並可同步建立活動品項與排班備註。",
                        "fields": "title、startsAt、endsAt、eventItems[{type/name/price/cost/servicePersonId/requiresInventory/inventoryQuantity/active}]。",
                        "rules": "活動品項 name 必填、price/cost 不可負；requiresInventory 時 inventoryQuantity 必須大於 0；會建立對應 Item，必要時新增 PURCHASE 庫存；會依活動日期建立 shift 備註。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/events/:id",
                        "content": "更新活動與活動品項。",
                        "fields": "title、startsAt、endsAt、eventItems[{id, type/name/price/cost/servicePersonId/requiresInventory/inventoryQuantity/active}]。",
                        "rules": "未提交的既有活動品項會被刪除；同步更新活動日期備註；新增需要庫存品項時補建立 PURCHASE 庫存。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/events/:id",
                        "content": "刪除活動。",
                        "fields": "路徑 id。",
                        "rules": "一般實體刪除；若關聯阻擋會回錯誤。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/employees",
                        "content": "活動品項服務人員下拉。",
                        "fields": "Employee。",
                        "rules": "活動建立品項時可指定服務人員。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/items",
                        "content": "活動品項同步後重新讀取品項，活動報名也用於品項下拉。",
                        "fields": "Item。",
                        "rules": "活動報名只能選該活動建立且 active=true 的品項。"
                    },
                    {
                        "method": "GET",
                        "api": "/api/event-registrations",
                        "content": "讀取活動報名資料。",
                        "fields": "id、eventId、itemId、registeredAt、attendeeName、phone、paidAmount、paymentMethod、remittanceRef、status、saleId、event、item、customer。",
                        "rules": "需 event-registrations 模組權限；此頁內嵌於 /events。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/event-registrations",
                        "content": "新增活動報名，付款時同步建立營收。",
                        "fields": "eventId、itemId、registeredAt、attendeeName、phone、paymentMethod、remittanceRef、isPaid。",
                        "rules": "itemId 必填且必須屬於該活動 active 品項；paidAmount 由品項 price 強制帶入；isPaid=true 時建立 Sale 與 SaleLine。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/event-registrations/:id",
                        "content": "更新活動報名或付款狀態。",
                        "fields": "eventId、itemId、registeredAt、attendeeName、phone、paymentMethod、remittanceRef、isPaid。",
                        "rules": "isPaid=false 時會刪除已同步的 Sale；isPaid=true 時新增或更新 Sale。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/event-registrations/:id",
                        "content": "刪除活動報名。",
                        "fields": "路徑 id。",
                        "rules": "若已同步 Sale，會先刪 SaleLine 與 Sale，再刪報名。"
                    }
                ]
            }
        ]
    },
    {
        "module": "服務",
        "pages": [
            {
                "page": "/sops",
                "name": "SOP",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/sops",
                        "content": "讀取 SOP 主檔。",
                        "fields": "id、title、category、status、description、createdAt、updatedAt。",
                        "rules": "需 sops 權限；依 createdAt desc 排序。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/sops",
                        "content": "新增 SOP。",
                        "fields": "title、category、status、description。",
                        "rules": "title 必填；status 預設 DRAFT，可為 DRAFT、ACTIVE、ARCHIVED。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/sops/:id",
                        "content": "更新 SOP。",
                        "fields": "title、category、status、description。",
                        "rules": "套用共用模組更新流程。"
                    },
                    {
                        "method": "DELETE",
                        "api": "/api/sops/:id",
                        "content": "刪除 SOP。",
                        "fields": "路徑 id。",
                        "rules": "一般實體刪除；若被 SOP 步驟引用可能回 409。"
                    }
                ]
            }
        ]
    },
    {
        "module": "系統",
        "pages": [
            {
                "page": "/permissions",
                "name": "權限設定",
                "apis": [
                    {
                        "method": "GET",
                        "api": "/api/permissions",
                        "content": "讀取頁面權限、目前使用者、角色與帳號。",
                        "fields": "pagePermissions、currentUser、roles、users。",
                        "rules": "需 permissions 權限；Admin 使用者回傳全部權限。"
                    },
                    {
                        "method": "POST",
                        "api": "/api/permissions",
                        "content": "新增角色或帳號。",
                        "fields": "type=role：name、notes、pagePermissions；type=user：email、password、roleId。",
                        "rules": "新增角色 name 必填；新增帳號 email/roleId 必填且密碼至少 6 碼；只能指派目前使用者可管理的權限。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/permissions/roles/:id",
                        "content": "更新角色名稱、備註、啟用狀態與權限。",
                        "fields": "name、notes、active、pagePermissions。",
                        "rules": "Admin 角色不可關閉；更新角色後會同步更新該角色下非 Admin 使用者的 pagePermissions。"
                    },
                    {
                        "method": "PATCH",
                        "api": "/api/permissions/users/:id",
                        "content": "更新帳號角色、姓名與啟用狀態。",
                        "fields": "name、active、roleId。",
                        "rules": "Admin 帳號保留全部權限，不可調整；role 必須存在且 active=true。"
                    }
                ]
            }
        ]
    }
]


METHOD_FILL = {
    "GET": "EAF4F7",
    "POST": "EAF5EA",
    "PATCH": "FFF3D6",
    "DELETE": "FCE8E6",
}

FIELD_LABELS = {
    "Employee": "人員物件",
    "Event": "活動物件",
    "Item": "進銷項目物件",
    "Sale": "營收物件",
    "Shift": "排班物件",
    "active": "啟用",
    "activityAnalysis": "活動分析",
    "afterAmount": "異動後金額",
    "amount": "金額",
    "api": "API",
    "attendeeName": "報名者姓名",
    "beforeAmount": "異動前金額",
    "category": "分類",
    "checkedOut": "出帳",
    "cost": "成本",
    "createdAt": "建立時間",
    "currentUser": "目前使用者",
    "customer": "客戶",
    "customerName": "客戶名稱",
    "date": "日期",
    "description": "說明",
    "discount": "折扣",
    "email": "Email",
    "employee": "員工",
    "employeeId": "員工 ID",
    "employeePayoutAnalysis": "員工薪資/分潤分析",
    "endDate": "結束日期",
    "endsAt": "結束時間",
    "event": "活動",
    "eventId": "活動 ID",
    "eventItems": "活動品項",
    "fixedExpenseAnalysis": "固定支出分析",
    "id": "識別碼",
    "inventoryQuantity": "庫存數量",
    "isPaid": "是否付款",
    "item": "品項",
    "itemAnalysis": "品項分析",
    "itemId": "品項 ID",
    "itemTypeAnalysis": "品項類型分析",
    "items": "品項",
    "lineTotal": "明細小計",
    "lines": "銷售明細",
    "location": "地點",
    "movedAt": "異動時間",
    "name": "名稱",
    "notes": "備註",
    "pagePermissions": "頁面權限",
    "password": "密碼",
    "paymentMethod": "付款方式",
    "paymentMethodAnalysis": "付款方式分析",
    "pettyCashExpenseAnalysis": "零用金支出分析",
    "phone": "電話",
    "price": "售價",
    "purpose": "用途",
    "quantity": "數量",
    "query.date": "查詢日期",
    "reason": "原因",
    "registeredAt": "報名日期",
    "remittanceRef": "匯款資訊",
    "requiresInventory": "需要庫存",
    "role": "角色",
    "roleId": "角色 ID",
    "roles": "角色清單",
    "saleId": "營收 ID",
    "salesPerson": "銷售人員",
    "salesPersonId": "銷售人員 ID",
    "serviceCommissionAnalysis": "服務抽成分析",
    "servicePerson": "服務人員",
    "servicePersonId": "服務人員 ID",
    "shiftIncomeAnalysis": "排班收入分析",
    "soldAt": "銷售日期",
    "startDate": "開始日期",
    "startsAt": "開始時間",
    "status": "狀態",
    "stockQuantity": "庫存數量",
    "subtotal": "小計",
    "summary": "摘要",
    "tagIds": "TAG ID",
    "tags": "TAG",
    "title": "標題",
    "total": "總計",
    "type": "類型",
    "unitPrice": "單價",
    "updatedAt": "更新時間",
    "users": "帳號清單",
}


def enrich_field_text(text: str) -> str:
    enriched = text
    for field_name in sorted(FIELD_LABELS, key=len, reverse=True):
        label = FIELD_LABELS[field_name]
        enriched = enriched.replace(f"{field_name}（", f"{field_name}（")
        import re

        pattern = rf"(?<![A-Za-z0-9_.]){re.escape(field_name)}(?![A-Za-z0-9_.])"
        enriched = re.sub(pattern, f"{field_name}（{label}）", enriched)
    return enriched


def set_font(run, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width(cell, width)


def set_cell_margins(table, top=90, bottom=90, start=120, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    mar = tbl_pr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def write_cell(cell, text: str, *, bold=False, size=8.2, color="000000", align=None, fill=None) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_font(run, size, bold, color)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 11.5, "1F4D78", 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_footer(doc: Document) -> None:
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("UNV.PLAN ERP API Reference")
    set_font(run, 8, False, "666666")


def add_title(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("ERP 頁面 API 詳細對照表")
    set_font(run, 22, True, "0B2545")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("層級：模塊 → 功能頁 → API。範圍依目前 localhost:3000 Dev UI 可見頁面與頁內分頁整理。")
    set_font(run, 9.5, False, "555555")


def add_api_table(doc: Document, apis: list[dict[str, str]]) -> None:
    headers = ["Method", "API", "內容", "欄位說明", "規則"]
    widths = [900, 2300, 2200, 2800, 2600]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_cell_margins(table)
    repeat_header(table.rows[0])
    for cell, header in zip(table.rows[0].cells, headers):
        write_cell(cell, header, bold=True, size=8.5, color="0B2545", align=WD_ALIGN_PARAGRAPH.CENTER, fill="E8EEF5")

    for api in apis:
        row = table.add_row()
        method = api["method"]
        cells = row.cells
        write_cell(cells[0], method, bold=True, size=8.0, align=WD_ALIGN_PARAGRAPH.CENTER, fill=METHOD_FILL.get(method, "FFFFFF"))
        write_cell(cells[1], api["api"], size=8.0)
        write_cell(cells[2], api["content"], size=8.0)
        write_cell(cells[3], enrich_field_text(api["fields"]), size=7.8)
        write_cell(cells[4], api["rules"], size=7.8)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx() -> None:
    doc = Document()
    configure_document(doc)
    add_footer(doc)
    add_title(doc)

    for section in SECTIONS:
        doc.add_heading(str(section["module"]), level=1)
        for page in section["pages"]:
            heading = doc.add_heading(f'{page["name"]}（{page["page"]}）', level=2)
            heading.keep_with_next = True
            add_api_table(doc, page["apis"])

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_docx()
